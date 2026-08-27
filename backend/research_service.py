import csv
import io
from datetime import datetime
from typing import Any, Iterable

from sqlalchemy import select
from sqlalchemy.orm import Session

from analytics_service import (
    assignment_catalog,
    participant_label as analytics_participant_label,
    record_assignment_open,
    student_analytics_payload,
)

from data_contracts import (
    AUTOMATIC_LEARNING_EVENT_TYPES,
    CURRENT_DATA_SCHEMA_VERSION,
    LearningEventType,
    normalize_event_data,
)
from models import LearningEvent, StudySession, User


AUTOMATIC_EVENT_TYPES = AUTOMATIC_LEARNING_EVENT_TYPES


def participant_label(user: User) -> str:
    """Return a research-safe identifier without exposing the login ID."""
    return analytics_participant_label(user)


def get_user_study_session(
    session: Session, study_session_id: str | None, user: User
) -> StudySession | None:
    if not study_session_id:
        return None
    return session.scalar(
        select(StudySession).where(
            StudySession.id == study_session_id,
            StudySession.user_id == user.id,
        )
    )


def record_learning_event(
    session: Session,
    *,
    study_session: StudySession,
    user: User,
    event_type: str,
    assignment_id: int | None = None,
    assignment_version: int | None = None,
    event_data: dict[str, Any] | None = None,
    client_occurred_at: datetime | None = None,
    commit: bool = True,
) -> LearningEvent:
    normalized_event_type = LearningEventType(str(event_type)).value
    if (assignment_id is None) != (assignment_version is None):
        raise ValueError(
            "assignment_id and assignment_version must be provided together"
        )
    normalized_event_data = normalize_event_data(
        normalized_event_type, event_data or {}
    )
    event = LearningEvent(
        study_session_id=study_session.id,
        user_id=user.id,
        assignment_id=assignment_id,
        assignment_version=assignment_version,
        event_type=normalized_event_type,
        event_data=normalized_event_data,
        data_schema_version=CURRENT_DATA_SCHEMA_VERSION,
        client_occurred_at=client_occurred_at,
    )
    session.add(event)
    if normalized_event_type == "assignment_opened":
        record_assignment_open(
            session,
            user=user,
            assignment_id=assignment_id,
            assignment_version=assignment_version,
        )
    if commit:
        session.commit()
        session.refresh(event)
    return event


def build_research_dataset(session: Session) -> dict[str, list[dict[str, Any]]]:
    students = session.scalars(
        select(User)
        .where(User.role == "student", User.is_demo.is_(False))
        .order_by(User.created_at)
    ).all()
    payloads = [student_analytics_payload(session, user) for user in students]
    summary_rows = [payload["summary"] for payload in payloads]
    attempt_rows = [row for payload in payloads for row in payload["attempts"]]
    event_rows = [row for payload in payloads for row in payload["events"]]
    assignment_rows = []
    for payload in payloads:
        participant_code = payload["summary"]["participant_code"]
        assignment_rows.extend(
            {"participant_code": participant_code, **row}
            for row in payload["assignment_summaries"]
        )
    return {
        "student_summary": summary_rows,
        "assignment_attempts": attempt_rows,
        "assignment_summary": assignment_rows,
        "class_assignment_summary": assignment_catalog(session, payloads),
        "learning_events": event_rows,
    }


def student_research_payload(session: Session, user: User) -> dict[str, Any]:
    return student_analytics_payload(session, user)


def csv_bytes(rows: Iterable[dict[str, Any]]) -> bytes:
    values = list(rows)
    if not values:
        return b""
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(values[0].keys()))
    writer.writeheader()
    for row in values:
        writer.writerow({key: _export_value(value) for key, value in row.items()})
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(dataset: dict[str, list[dict[str, Any]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in dataset.items():
        sheet = workbook.create_sheet(_sheet_title(sheet_name))
        if not rows:
            sheet.append(["No records"])
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C6B52")
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append([_export_value(row.get(header)) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(36, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _export_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        import json

        return json.dumps(value, sort_keys=True, separators=(",", ":"))
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _sheet_title(value: str) -> str:
    return " ".join(word.capitalize() for word in value.split("_"))[:31]
