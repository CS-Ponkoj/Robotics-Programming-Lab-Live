"""Build student-safe web pages from the six source content folders.

Run with a Python environment containing python-docx 1.2.0. Generated pages
are checked into backend/course_content; python-docx is not needed at runtime.
"""

from __future__ import annotations

import hashlib
import html
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "course_materials" / "content"
OUTPUT = ROOT / "backend" / "course_content"
ASSETS = OUTPUT / "assets"

TOPICS = {
    "Obstacle Avoidance": (
        "obstacle-avoidance",
        "Obstacle Avoidance",
        "Detect obstacles and choose safe routes around them.",
    ),
    "Path Planning": (
        "path-planning",
        "Path Planning",
        "Explore search algorithms and robot routes through constrained spaces.",
    ),
    "Sensor Fusion": (
        "sensor-fusion",
        "Sensor Fusion",
        "Combine readings from multiple sensors to estimate robot state.",
    ),
    "Range Sensors": (
        "range-sensors",
        "Range Sensors",
        "Compare distance sensors and interpret obstacle measurements.",
    ),
    "Locomotion": (
        "locomotion",
        "Locomotion",
        "Study robot mobility, stability, and motion calculations.",
    ),
    "Hybrid Navigation GPS": (
        "hybrid-navigation-gps",
        "Hybrid Navigation and GPS",
        "Navigate when satellite positioning is unreliable or unavailable.",
    ),
}

# Practice exercises whose graded, possibly simplified version is an app assignment.
GRADED_ASSIGNMENTS = {
    ("obstacle-avoidance", "assignment-bug-algorithm-scenario-1"): 1,
    ("obstacle-avoidance", "assignment-bug-algorithm-scenario-2"): 2,
    ("obstacle-avoidance", "assignment-oa"): 3,
    ("obstacle-avoidance", "assignment-vfh"): 4,
}

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
ANSWER_START = re.compile(
    r"^(?:(?:answer|solution|solutions)\s*(?::|$)|solution outline\b|instructor notes?\b)",
    re.I,
)
QUESTION_START = re.compile(r"^question(?:\s+\d+)?\s*(?::|$)", re.I)
CHOICE = re.compile(r"^([A-D])[.)]\s*(.+)", re.I)
MATH_NAMESPACES = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
WORD_NAMESPACES = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def student_source(path: Path) -> bool:
    parts = [part.lower() for part in path.relative_to(SOURCE).parts]
    return not any("reference" in part for part in parts) and not any(
        marker in path.stem.lower() for marker in ("solution", "unittest", "step_by_step_explanation")
    )


def visible_paragraphs(document: Document, *, strip_answers: bool) -> set:
    visible: set = set()
    in_answer = False
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if strip_answers and ANSWER_START.match(text):
            in_answer = True
            continue
        if strip_answers and (QUESTION_START.match(text) or text.lower().startswith("extra question:")):
            in_answer = False
        if strip_answers and text.lower().startswith("references:"):
            in_answer = True
        if not in_answer:
            visible.add(paragraph._p)
    return visible


def image_names(document: Document, paragraph: Paragraph) -> list[str]:
    names: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relation = blip.get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        if not relation or relation not in document.part.related_parts:
            continue
        part = document.part.related_parts[relation]
        extension = Path(str(part.partname)).suffix.lower()
        if extension not in IMAGE_TYPES:
            continue
        blob = part.blob
        name = f"{hashlib.sha256(blob).hexdigest()[:20]}{extension}"
        destination = ASSETS / name
        if not destination.exists():
            destination.write_bytes(blob)
        names.append(name)
    return names


def figure_html(names: list[str], alt: str, caption: str = "") -> str:
    images = "".join(
        f'<img src="/content-assets/{name}" alt="{html.escape(alt, quote=True)}" '
        'loading="lazy" />'
        for name in names
    )
    caption_html = f"<figcaption>{caption}</figcaption>" if caption else ""
    return f"<figure>{images}{caption_html}</figure>"


def local_name(element) -> str:
    return element.tag.rsplit("}", 1)[-1]


def omml_html(element) -> str:
    """Render Word equation markup as inline HTML, keeping every character."""
    name = local_name(element)
    if name == "t":
        return html.escape(element.text or "")
    if name.endswith("Pr"):
        return ""
    parts = {local_name(child): child for child in element}

    def part(key: str) -> str:
        return omml_html(parts[key]) if key in parts else ""

    def property_value(properties: str, key: str, default: str) -> str:
        found = element.find(f"m:{properties}/m:{key}", MATH_NAMESPACES)
        return found.get(f"{{{MATH_NAMESPACES['m']}}}val", default) if found is not None else default

    if name == "f":
        return f"<sup>{part('num')}</sup>⁄<sub>{part('den')}</sub>"
    if name == "sSup":
        return f"{part('e')}<sup>{part('sup')}</sup>"
    if name == "sSub":
        return f"{part('e')}<sub>{part('sub')}</sub>"
    if name == "sSubSup":
        return f"{part('e')}<sub>{part('sub')}</sub><sup>{part('sup')}</sup>"
    if name == "rad":
        degree = part("deg")
        return f"{f'<sup>{degree}</sup>' if degree else ''}√({part('e')})"
    if name == "d":
        begin = property_value("dPr", "begChr", "(")
        end = property_value("dPr", "endChr", ")")
        separator = property_value("dPr", "sepChr", ",")
        items = [omml_html(child) for child in element if local_name(child) == "e"]
        return f"{html.escape(begin)}{html.escape(separator).join(items)}{html.escape(end)}"
    if name == "nary":
        symbol = property_value("naryPr", "chr", "∫")
        limits = "".join(
            f"<{tag}>{part(key)}</{tag}>" for key, tag in (("sub", "sub"), ("sup", "sup"))
            if part(key)
        )
        return f"{html.escape(symbol)}{limits}{part('e')}"
    return "".join(omml_html(child) for child in element)


def inline_html(paragraph: Paragraph) -> str:
    """Paragraph text with links, emphasis, sub/superscripts, and equations in place."""
    runs: list[str] = []
    for child in paragraph._p:
        name = local_name(child)
        if name == "hyperlink":
            item = Hyperlink(child, paragraph)
            text = html.escape(item.text)
            url = item.url or ""
            if url.startswith(("https://", "http://")):
                runs.append(f'<a href="{html.escape(url, quote=True)}">{text}</a>')
            else:
                runs.append(text)
        elif name == "r" and child.tag == qn("w:r"):
            run = Run(child, paragraph)
            text = html.escape(run.text).replace("\n", "<br />")
            if not text:
                continue
            if run.font.subscript:
                text = f"<sub>{text}</sub>"
            elif run.font.superscript:
                text = f"<sup>{text}</sup>"
            if run.bold:
                text = f"<strong>{text}</strong>"
            if run.italic:
                text = f"<em>{text}</em>"
            runs.append(text)
        elif name in {"oMath", "oMathPara"}:
            equations = [child] if name == "oMath" else child.findall("m:oMath", MATH_NAMESPACES)
            for equation in equations:
                runs.append(f'<span class="math-expression">{omml_html(equation)}</span>')
    text_html = "".join(runs)
    if not text_html and paragraph.text.strip():
        text_html = html.escape(paragraph.text.strip())
    return text_html


def numbering_properties(paragraph: Paragraph):
    """The paragraph's Word numbering, set directly or through its style."""
    for owner in (paragraph._p, paragraph.style.element if paragraph.style else None):
        properties = owner.pPr if owner is not None else None
        if properties is not None and properties.numPr is not None:
            return properties.numPr
    return None


def list_details(document: Document, paragraph: Paragraph) -> tuple[str, int, str, int, str] | None:
    """(tag, level, type attribute, first number, numbering id) for list paragraphs."""
    style = paragraph.style.name.lower() if paragraph.style else ""
    properties = numbering_properties(paragraph)
    number_id = properties.numId.val if properties is not None and properties.numId is not None else 0
    if not number_id:
        if "list bullet" in style:
            return "ul", 0, "", 1, "style-bullet"
        if "list number" in style:
            return "ol", 0, "", 1, "style-number"
        return None
    level = properties.ilvl.val if properties.ilvl is not None else 0
    number_format, start = "bullet", 1
    try:
        numbering = document.part.numbering_part.element
    except (KeyError, NotImplementedError):
        numbering = None
    if numbering is not None:
        instance = numbering.find(f'w:num[@w:numId="{number_id}"]', WORD_NAMESPACES)
        if instance is not None:
            abstract_id = instance.find("w:abstractNumId", WORD_NAMESPACES).get(qn("w:val"))
            definition = numbering.find(
                f'w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]',
                WORD_NAMESPACES,
            )
            if definition is not None:
                found_format = definition.find("w:numFmt", WORD_NAMESPACES)
                found_start = definition.find("w:start", WORD_NAMESPACES)
                if found_format is not None:
                    number_format = found_format.get(qn("w:val"))
                if found_start is not None:
                    start = int(found_start.get(qn("w:val")))
    ordered = {"decimal": "", "lowerLetter": "a", "upperLetter": "A",
               "lowerRoman": "i", "upperRoman": "I"}
    if number_format in ordered:
        return "ol", level, ordered[number_format], start, str(number_id)
    return "ul", level, "", 1, str(number_id)


def heading_level(paragraph: Paragraph, styled_document: bool) -> int | None:
    """Heading level from a Word heading style, or from a short bold-only line."""
    style = paragraph.style.name.lower() if paragraph.style else ""
    if style == "title":
        return 2
    if "heading" in style:
        match = re.search(r"(\d+)", style)
        return min(4, int(match.group(1)) + 1) if match else 2
    if paragraph._p.xpath("./w:hyperlink | .//m:oMath | .//a:blip"):
        return None
    runs = [run for run in paragraph.runs if run.text.strip()]
    text = paragraph.text.strip()
    if not runs or not all(run.bold for run in runs):
        return None
    if len(text) > 90 or text.endswith(".") or any(mark in text for mark in "=\\|"):
        return None
    sizes = [run.font.size.pt for run in runs if run.font.size]
    size = max(sizes) if sizes else 0
    if size >= 15:
        return 2
    if size >= 13:
        return 3
    if styled_document:
        return 4
    return 4 if text.endswith(":") else 3


def split_blank_lines(text_html: str) -> list[str]:
    """Split a paragraph at blank lines typed inside it, when no formatting spans the break."""
    parts = [part.strip() for part in re.split(r"(?:<br />\s*){2,}", text_html)]
    if len(parts) == 1 or any(
        part.count(f"<{tag}") != part.count(f"</{tag}>")
        for part in parts for tag in ("strong", "em", "a", "sub", "sup", "span")
    ):
        return [text_html]
    return [part for part in parts if part]


CAPTION = re.compile(r"^(?:fig\.?|figure|image from)\b", re.I)
CAPTION_PREFIX = re.compile(r"^(?:figure|fig\.?)\s*\d*\s*[:.]?\s*", re.I)


def table_html(table: Table) -> str:
    rows = [[html.escape(cell.text).replace("\n", "<br />") for cell in row.cells]
            for row in table.rows]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    # An empty top-left cell means the first column labels each row.
    row_labels = not header[0].strip()
    head = "".join(f'<th scope="col">{cell}</th>' for cell in header)
    body_rows = []
    for row in body:
        cells = [
            f'<th scope="row">{cell}</th>' if index == 0 and row_labels else f"<td>{cell}</td>"
            for index, cell in enumerate(row)
        ]
        body_rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"


def nest_lists(blocks: list[tuple]) -> list[str]:
    """Join consecutive list items into nested HTML lists."""
    output: list[str] = []
    stack: list[tuple[str, int]] = []
    counters: dict[tuple[str, int], int] = {}

    def close_above(level: int) -> None:
        while stack and stack[-1][1] > level:
            output.append(f"</li></{stack.pop()[0]}>")

    for block in blocks:
        if block[0] != "item":
            close_above(-1)
            output.append(block[1])
            continue
        _, content, (tag, level, kind, start, number_id) = block
        # Word continues numbering until a shallower item restarts deeper levels.
        for key in [key for key in counters if key[0] == number_id and key[1] > level]:
            del counters[key]
        number = counters.get((number_id, level), start - 1) + 1
        counters[(number_id, level)] = number
        close_above(level)
        if stack and stack[-1][1] == level and stack[-1][0] != tag:
            output.append(f"</li></{stack.pop()[0]}>")
        if stack and stack[-1][1] == level:
            output.append("</li>")
        else:
            attributes = f' type="{kind}"' if kind else ""
            if tag == "ol" and number != 1:
                attributes += f' start="{number}"'
            output.append(f"<{tag}{attributes}>")
            stack.append((tag, level))
        output.append(f"<li>{content}")
    close_above(-1)
    return output


def document_html(path: Path, *, strip_answers: bool = False) -> str:
    document = Document(path)
    visible = visible_paragraphs(document, strip_answers=strip_answers)
    label = re.sub(r"\s+", " ", path.stem.replace("_", " ")).strip()
    styled_document = any(
        paragraph.style is not None
        and ("heading" in paragraph.style.name.lower() or paragraph.style.name.lower() == "title")
        for paragraph in document.paragraphs
    )
    content = [
        block for block in document.iter_inner_content()
        if isinstance(block, Table) or block._p in visible
    ]
    blocks: list[tuple] = []
    index = 0
    while index < len(content):
        block = content[index]
        index += 1
        if isinstance(block, Table):
            blocks.append(("block", table_html(block)))
            continue
        text_html = inline_html(block)
        images = image_names(document, block)
        if images and not text_html.strip():
            # The next "Figure ..." line, past blank paragraphs, becomes the caption.
            ahead = index
            while (ahead < len(content) and isinstance(content[ahead], Paragraph)
                   and not content[ahead].text.strip()
                   and not image_names(document, content[ahead])):
                ahead += 1
            following = content[ahead] if ahead < len(content) else None
            if (isinstance(following, Paragraph) and CAPTION.match(following.text.strip())
                    and not image_names(document, following)):
                caption = following.text.strip()
                alt = CAPTION_PREFIX.sub("", caption).strip() or f"Diagram from {label}"
                if alt.lower() == "source":
                    alt = f"Diagram from {label}"
                blocks.append(("block", figure_html(images, alt, inline_html(following))))
                index = ahead + 1
            else:
                blocks.append(("block", figure_html(images, f"Diagram from {label}")))
            continue
        figures = figure_html(images, f"Diagram from {label}") if images else ""
        if not text_html.strip():
            continue
        details = list_details(document, block)
        if details is not None:
            blocks.append(("item", text_html + figures, details))
            continue
        level = heading_level(block, styled_document)
        if level is not None:
            blocks.append(("block", f"<h{level}>{text_html}</h{level}>{figures}"))
        else:
            paragraphs = "".join(f"<p>{part}</p>" for part in split_blank_lines(text_html))
            blocks.append(("block", f"{paragraphs}{figures}"))
    return "\n".join(part for part in nest_lists(blocks) if part)


def text_html(path: Path) -> str:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return "\n".join(
        f"<p>{html.escape(line.strip())}</p>"
        for line in text.splitlines()
        if line.strip()
    )


def quiz_questions(path: Path) -> tuple[list[dict], dict[str, str]]:
    document = Document(path)
    questions: list[dict] = []
    answers: dict[str, str] = {}
    prompt: list[str] = []
    options: list[dict] = []
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if not value:
            continue
        answer = re.match(r"^Answer\s*:\s*([A-D])\b", value, re.I)
        if answer:
            question_id = f"q{len(questions) + 1}"
            if prompt and len(options) >= 2:
                questions.append({
                    "id": question_id,
                    "prompt": "\n".join(prompt),
                    "options": options,
                })
                answers[question_id] = answer.group(1).upper()
            prompt, options = [], []
            continue
        for line in value.splitlines():
            choice = CHOICE.match(line.strip())
            if choice:
                options.append({"id": choice.group(1).upper(), "text": choice.group(2)})
            elif line.strip():
                prompt.append(line.strip())
    if prompt or options or len(questions) != len(answers):
        raise ValueError(f"Incomplete answer key in {path}")
    return questions, answers


def label_for_group(path: Path, topic_root: Path) -> str:
    parts = path.relative_to(topic_root).parts
    readable = [part.replace("_", " ").replace("-", " ") for part in parts]
    readable = [re.sub(r"^Assignment\s+", "", value, flags=re.I) for value in readable]
    return " · ".join(readable).replace("  ", " ")


def build() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    catalog: list[dict] = []
    answer_key: dict[str, dict[str, str]] = {}
    authored_quizzes = json.loads(
        (ROOT / "scripts" / "course_quiz_bank.json").read_text(encoding="utf-8")
    )
    for folder_name, (topic_id, title, summary) in TOPICS.items():
        topic_root = SOURCE / folder_name
        topic_dir = OUTPUT / topic_id
        direct_docs = sorted(
            path for path in topic_root.glob("*.docx") if student_source(path)
        )
        lessons = [
            path for path in direct_docs
            if "quiz" not in path.stem.lower() and "math" not in path.stem.lower()
        ]
        if len(lessons) != 1:
            raise ValueError(f"Expected one lesson document for {folder_name}: {lessons}")
        sections: list[dict] = []

        def add_section(section_id: str, name: str, kind: str, body: str,
                        questions: list[dict] | None = None) -> None:
            metadata = {"id": section_id, "title": name, "kind": kind}
            sections.append(metadata)
            write_json(topic_dir / f"{section_id}.json", {
                **metadata, "topic_id": topic_id, "html": body,
                "questions": questions or [],
            })

        add_section("lesson", "Lesson", "lesson", document_html(lessons[0]))

        exercise_files: dict[Path, list[Path]] = defaultdict(list)
        for path in sorted(topic_root.rglob("*")):
            if not path.is_file() or path.parent == topic_root or not student_source(path):
                continue
            if path.suffix.lower() in {".docx", ".txt"}:
                exercise_files[path.parent].append(path)
        for directory, files in sorted(exercise_files.items()):
            section_id = slug(str(directory.relative_to(topic_root)))
            name = label_for_group(directory, topic_root)
            pieces = []
            assignment_number = GRADED_ASSIGNMENTS.get((topic_id, section_id))
            if assignment_number is not None:
                pieces.append(
                    "<p><strong>Note:</strong> This is the original practice version of "
                    f"Assignment {assignment_number}. The graded version on the Assignments "
                    "page may use different values or a simplified model. Where they differ, "
                    "follow the Assignments page.</p>"
                )
            ordered = sorted(files, key=lambda path: (
                "instruction" not in path.stem.lower()
                and "scenario" not in path.stem.lower()
                and "senario" not in path.stem.lower(),
                path.name.lower(),
            ))
            for path in ordered:
                if path.suffix.lower() == ".docx":
                    pieces.append(document_html(path, strip_answers=True))
                elif path.suffix.lower() == ".txt":
                    pieces.append(text_html(path))
                else:
                    pieces.append(text_html(path))
            add_section(section_id, name, "exercise", "\n".join(pieces))

        source_quiz = next(
            (path for path in direct_docs if path.stem.lower() == "quiz"), None
        )
        if source_quiz is not None:
            questions, answers = quiz_questions(source_quiz)
        else:
            rows = authored_quizzes[topic_id]
            questions, answers = [], {}
            for number, (prompt, choices, correct) in enumerate(rows, start=1):
                if len(choices) != 4 or correct not in "ABCD":
                    raise ValueError(f"Invalid authored question in {topic_id}")
                # Rotate the location of the right answer; it must not always be A.
                rotation = (number - 1) % 4
                rotated = choices[rotation:] + choices[:rotation]
                question_id = f"q{number}"
                questions.append({
                    "id": question_id,
                    "prompt": prompt,
                    "options": [
                        {"id": letter, "text": choice}
                        for letter, choice in zip("ABCD", rotated)
                    ],
                })
                answers[question_id] = "ABCD"[("ABCD".index(correct) - rotation) % 4]
        if len(questions) < 10 or len(answers) != len(questions):
            raise ValueError(f"Expected at least 10 complete quiz questions for {topic_id}")
        section_id = "quiz"
        version = hashlib.sha256(
            json.dumps([questions, answers], ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()[:16]
        add_section(section_id, "Quiz", "quiz", "", questions)
        quiz_file = topic_dir / "quiz.json"
        quiz_payload = json.loads(quiz_file.read_text(encoding="utf-8"))
        quiz_payload["version"] = version
        write_json(quiz_file, quiz_payload)
        answer_key[f"{topic_id}/{section_id}"] = answers

        for path in direct_docs:
            if path == lessons[0] or "quiz" in path.stem.lower():
                continue
            add_section(
                slug(path.stem), path.stem.replace("_", " "),
                "math", document_html(path, strip_answers=True),
            )

        write_json(topic_dir / "index.json", {
            "id": topic_id, "title": title, "summary": summary,
            "sections": sections,
        })
        # Remove pages from a previous build that are no longer listed.
        current = {"index.json", *(f"{section['id']}.json" for section in sections)}
        for old_page in topic_dir.glob("*.json"):
            if old_page.name not in current:
                old_page.unlink()
        catalog.append({
            "id": topic_id, "title": title, "summary": summary,
            "section_counts": {
                kind: sum(section["kind"] == kind for section in sections)
                for kind in ("lesson", "exercise", "quiz", "math")
            },
        })
    write_json(OUTPUT / "catalog.json", {"topics": catalog})
    write_json(OUTPUT / "quiz_answers.json", answer_key)
    print(f"Built {len(catalog)} topics and {sum(len(v) for v in answer_key.values())} quiz questions")


if __name__ == "__main__":
    build()
